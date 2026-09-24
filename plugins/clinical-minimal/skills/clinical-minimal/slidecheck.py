"""Per-slide visual verification for Clinical Minimal decks.

    python slidecheck.py deck.pptx                    render + automatic checks + review status
    python slidecheck.py deck.pptx --mark 1-4,7 pass  record your visual review of those slides
    python slidecheck.py deck.pptx --mark 5 fix "arrow crosses label"

PowerPoint lays the deck out (slideprobe.ps1), so text sizes, line breaks and overflow are the real
ones, not estimates. Writes <deck dir>/_check/<deck name>/: slide-NN.png, slide-NN.flag.png (issues boxed:
red = FAIL, amber = WARN), sheet.png (all slides), report.md, review.json.
A review is tied to the slide's current layout: change the slide and it must be reviewed again.
Exit code 0 only when there is no FAIL and every slide passed visual review.
"""
from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).parent
PX = 1600 / 960  # PNG pixels per point (slides are 960 x 540 pt)
PICTURE, CHART, TABLE, LINE, TEXTBOX, AUTOSHAPE = 13, 3, 19, 9, 17, 1
CHECKLIST = """Visual review: Read every slide-NN.png at full size and check, slide by slide:
1. Nothing clipped, cut off or colliding: text, arrows through labels, legends over bars, footer overlap.
2. Every label readable at a distance; no word split across lines; no one-word last line in a title.
3. Picture shows this slide's specific subject, is sharp, its subject (faces, labels) is not cropped away,
   no watermark or logo.
4. Balance: the visual fills the content area; no big empty block, no crowded corner.
5. The visual proves the title's claim (right numbers, right highlight colour, right direction).
6. Consistent with the other slides: title position, source line, credit position, build order.
Then: slidecheck.py deck.pptx --mark <slides> pass   (or fix "what is wrong", fix it, re-run)."""


# ---------------------------------------------------------------- probe
def probe(deck: Path, out: Path) -> dict:
    subprocess.run(["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
                    str(HERE / "slideprobe.ps1"), str(deck), str(out)], check=True)
    return json.loads((out / "probe.json").read_text(encoding="utf-8-sig"))


def _as_list(v):
    return v if isinstance(v, list) else ([] if v is None else [v])


def _rect(sh, bound=False):
    x, y, w, h = sh["bound"] if bound else (sh["x"], sh["y"], sh["w"], sh["h"])
    return x, y, x + w, y + h


def _overlap(a, b, tol=2.0):
    return min(a[2], b[2]) - max(a[0], b[0]) > tol and min(a[3], b[3]) - max(a[1], b[1]) > tol


def _is_title(sh):
    """Content-slide takeaway title: 26–32pt at the top of the slide."""
    return 26 <= max(_as_list(sh.get("sizes")) or [0]) <= 32 and sh["y"] < 60


def _is_small_print(sh, H):
    t = sh["text"].strip()
    return sh["y"] > H - 70 or t.startswith(("Image", "Source", "Images")) or t.isdigit()


# ---------------------------------------------------------------- checks
def check_slide(s, W, H, pics):
    """Return [(severity, rule, message, rect_pt)] for one probed slide."""
    issues = []
    shapes = s["shapes"]
    texts = [sh for sh in shapes if sh.get("bound")]
    for sh in texts:
        sh["lines"] = _as_list(sh.get("lines")); sh["sizes"] = _as_list(sh.get("sizes"))
    add = lambda sev, rule, msg, r: issues.append((sev, rule, msg, r))
    snippet = lambda sh: sh["text"].strip().replace("\r", " / ")[:50]

    for sh in texts:
        box, tb = _rect(sh), _rect(sh, bound=True)
        big = max(sh["sizes"] or [0])
        # 1-2 overflow: the text needs more room than its box
        if tb[3] > box[3] + 3 or tb[1] < box[1] - 3:
            add("FAIL", "overflow", f"text taller than its box ({tb[3] - tb[1]:.0f} > {sh['h']:.0f} pt): "
                f"'{snippet(sh)}'", tb)
        if tb[2] - tb[0] > sh["w"] + 3:
            add("FAIL", "overflow", f"text wider than its box: '{snippet(sh)}'", tb)
        # 3 off the slide
        if tb[0] < -1 or tb[1] < -1 or tb[2] > W + 1 or tb[3] > H + 1:
            add("FAIL", "off-slide", f"text runs off the slide: '{snippet(sh)}'", tb)
        # 4 word split across lines
        for a, b in zip(sh["lines"], sh["lines"][1:]):
            if a and b and a[-1].isalnum() and b[0].isalnum():
                add("FAIL", "word-split", f"word broken across lines: '…{a[-12:]}|{b[:12]}…'", tb)
                break
        # 5-6 title length and widows
        n = len([ln for ln in sh["lines"] if ln.strip()])
        if _is_title(sh) and n > 1:
            add("FAIL", "title-wrap", f"title wraps to {n} lines; shorten it: '{snippet(sh)}'", tb)
        if big >= 40 and n > 3:
            add("WARN", "title-wrap", f"display title runs {n} lines: '{snippet(sh)}'", tb)
        wrapped_last = n > 1 and not sh["lines"][-2].endswith(("\r", "\n", "\x0b"))  # not a paragraph end
        if big >= 20 and wrapped_last and len(sh["lines"][-1].split()) == 1:
            add("WARN", "widow", f"last line is one word: '{sh['lines'][-1].strip()}'", tb)
        # 7 size floor
        small = min(sh["sizes"] or [99])
        if small < 9:
            add("FAIL", "font-size", f"{small:g}pt text: '{snippet(sh)}'", tb)
        elif small < 12 and not _is_small_print(sh, H):
            add("WARN", "font-size", f"{small:g}pt label (minimum 12): '{snippet(sh)}'", tb)
        # source line must stay on one line so it does not climb into the content
        if sh["y"] > H - 70 and n > 1 and big <= 10:
            add("FAIL", "footer", "source/footer line wraps to two lines; shorten it or split the source", tb)

    # 8 text colliding with text
    for i, a in enumerate(texts):
        for b in texts[i + 1:]:
            if _overlap(_rect(a, True), _rect(b, True)):
                ra, rb = _rect(a, True), _rect(b, True)
                add("FAIL", "collision", f"text overlaps text: '{snippet(a)[:25]}' × '{snippet(b)[:25]}'",
                    (min(ra[0], rb[0]), min(ra[1], rb[1]), max(ra[2], rb[2]), max(ra[3], rb[3])))
    # 9-10 text over pictures, charts, tables; lines through text
    blocks = [sh for sh in shapes if sh["type"] in (PICTURE, CHART, TABLE) or sh.get("rows")]
    for bl in blocks:
        for t in texts:
            if t["id"] != bl["id"] and _overlap(_rect(bl), _rect(t, True), 4):
                add("WARN", "covered", f"text sits on a picture/chart/table: '{snippet(t)[:30]}'", _rect(t, True))
    for ln in (sh for sh in shapes if sh["type"] == LINE and min(sh["w"], sh["h"]) < 2):
        for t in texts:
            if _overlap(_rect(ln), _rect(t, True), 1):
                add("WARN", "line-through", f"line/arrow crosses text: '{snippet(t)[:30]}'", _rect(t, True))

    # 11-12 pictures: sharpness, distortion, credit
    has_credit = any(t["text"].strip().startswith("Image") for t in texts)
    for sh in (sh for sh in shapes if sh["type"] == PICTURE):
        p = pics.get(sh["id"])
        if p:
            cw, ch = p["px"][0] * (1 - p["crop"][0] - p["crop"][2]), p["px"][1] * (1 - p["crop"][1] - p["crop"][3])
            ppi = cw / (sh["w"] / 72)
            if ppi < 72:
                add("FAIL", "blurry", f"picture only {ppi:.0f} ppi at this size (need 110+)", _rect(sh))
            elif ppi < 110:
                add("WARN", "blurry", f"picture {ppi:.0f} ppi at this size (110+ looks sharp)", _rect(sh))
            if abs((cw / ch) / (sh["w"] / sh["h"]) - 1) > 0.03:
                add("FAIL", "stretched", "picture aspect ratio distorted", _rect(sh))
        if sh["w"] * sh["h"] > 2 * 72 * 72 and not has_credit:
            add("FAIL", "credit", "picture without an 'Image: <site>' credit", _rect(sh))

    # 13-14 content slides: a visual and builds
    title = next((t for t in texts if _is_title(t)), None)
    if title and not title["text"].strip().lower().startswith("reference"):
        drawn = [sh for sh in shapes if sh["type"] in (AUTOSHAPE, LINE) or sh["type"] in (PICTURE, CHART, TABLE)]
        if not blocks and len(drawn) < 3:
            add("FAIL", "no-visual", "content slide has no chart, diagram, table or picture", (36, 110, W - 36, H - 60))
        if not s["builds"]:
            add("WARN", "no-builds", "content slide has no click builds", (36, 20, 200, 60))
    if not s["transition"]:
        add("WARN", "no-transition", "slide has no transition", (W - 200, 20, W - 36, 60))
    return issues


def blank_share(png: Path) -> float:
    """Share of the content area (y 1.6–6.6 in, inside the margins) that is empty white."""
    from PIL import Image, ImageStat
    im = Image.open(png).convert("L")
    x0, y0, x1, y1 = int(36 * PX), int(115 * PX), int(924 * PX), int(475 * PX)
    cols, rows, empty = 12, 6, 0
    cw, ch = (x1 - x0) / cols, (y1 - y0) / rows
    for r in range(rows):
        for c in range(cols):
            st = ImageStat.Stat(im.crop((int(x0 + c * cw), int(y0 + r * ch), int(x0 + (c + 1) * cw), int(y0 + (r + 1) * ch))))
            empty += st.mean[0] > 250 and st.stddev[0] < 4
    return empty / (rows * cols)


def picture_info(deck: Path) -> list[dict]:
    """Per slide {shape id: native pixel size and crop fractions} from the pptx itself."""
    from pptx import Presentation
    out = []
    for slide in Presentation(str(deck)).slides:
        d = {}
        for sh in slide.shapes:
            if sh.shape_type == PICTURE:
                d[sh.shape_id] = {"px": sh.image.size,
                                  "crop": (sh.crop_left, sh.crop_top, sh.crop_right, sh.crop_bottom)}
        out.append(d)
    return out


# ---------------------------------------------------------------- outputs
def annotate(png: Path, issues, dst: Path):
    from PIL import Image, ImageDraw, ImageFont
    im = Image.open(png).convert("RGB")
    dr = ImageDraw.Draw(im)
    font = ImageFont.truetype("C:/Windows/Fonts/LeelaUIb.ttf", 22)
    for i, (sev, _, _, r) in enumerate(issues, 1):
        col = (200, 30, 30) if sev == "FAIL" else (220, 140, 0)
        box = [r[0] * PX - 4, r[1] * PX - 4, r[2] * PX + 4, r[3] * PX + 4]
        dr.rectangle(box, outline=col, width=4)
        dr.rectangle([box[0], box[1] - 28, box[0] + 34, box[1]], fill=col)
        dr.text((box[0] + 8, box[1] - 28), str(i), fill="white", font=font)
    im.save(dst)


def contact_sheet(out: Path, slides, status):
    from PIL import Image, ImageDraw, ImageFont
    tw, th, cols, pad = 480, 270, 4, 16
    rows = (len(slides) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * (tw + pad) + pad, rows * (th + 40 + pad) + pad), "white")
    dr = ImageDraw.Draw(sheet)
    font = ImageFont.truetype("C:/Windows/Fonts/LeelaUIb.ttf", 20)
    for i, s in enumerate(slides):
        x, y = pad + (i % cols) * (tw + pad), pad + (i // cols) * (th + 40 + pad)
        im = Image.open(out / s["png"]).convert("RGB").resize((tw, th))
        sheet.paste(im, (x, y + 34))
        dr.rectangle([x, y + 34, x + tw - 1, y + 34 + th - 1], outline=(200, 205, 210))
        label, col = status[s["index"]]
        dr.text((x, y + 4), f"{s['index']:02d}  {label}", fill=col, font=font)
    sheet.save(out / "sheet.png")


def slide_hash(s) -> str:
    keep = [{k: sh.get(k) for k in ("id", "type", "x", "y", "w", "h", "text", "lines")} for sh in s["shapes"]]
    return hashlib.sha1(json.dumps(keep, sort_keys=True).encode()).hexdigest()[:12]


def parse_range(spec: str) -> list[int]:
    out = []
    for part in spec.split(","):
        a, _, b = part.partition("-")
        out += list(range(int(a), int(b or a) + 1))
    return out


# ---------------------------------------------------------------- main
def run(deck: Path) -> int:
    out = deck.parent / "_check" / deck.stem
    for old in out.glob("slide-*.png"):
        old.unlink()
    data = probe(deck, out)
    W, H = data["width"], data["height"]
    pics = picture_info(deck)
    review_path = out / "review.json"
    review = json.loads(review_path.read_text("utf-8")) if review_path.exists() else {}
    slides = _as_list(data["slides"])
    lines, status, fails, unreviewed = [f"# Slide check: {deck.name}", ""], {}, 0, []
    for s in slides:
        s["shapes"] = _as_list(s["shapes"])
        issues = check_slide(s, W, H, pics[s["index"] - 1])
        title = next((sh for sh in s["shapes"] if sh.get("bound") and _is_title(sh)), None)
        if title and not title["text"].strip().lower().startswith("reference") and blank_share(out / s["png"]) > 0.5:
            issues.append(("WARN", "sparse", "over half the content area is empty", (36, 115, W - 36, 475)))
        n_fail = sum(sev == "FAIL" for sev, *_ in issues)
        fails += n_fail
        h = slide_hash(s)
        rv = review.get(str(s["index"]), {})
        reviewed = rv.get("hash") == h
        if not reviewed or rv.get("verdict") != "pass":
            unreviewed.append(s["index"])
        if issues:
            annotate(out / s["png"], issues, out / s["png"].replace(".png", ".flag.png"))
        verdict = rv.get("verdict") if reviewed else "not reviewed"
        label = f"{n_fail} fail, {len(issues) - n_fail} warn · {verdict}" if issues else f"clean · {verdict}"
        status[s["index"]] = (label, (180, 35, 24) if n_fail or verdict == "fix" else
                              (30, 123, 69) if verdict == "pass" else (74, 85, 97))
        review.setdefault(str(s["index"]), {})["current"] = h
        lines.append(f"## Slide {s['index']}: {label}")
        lines += [f"{i}. **{sev}** {rule}: {msg}" for i, (sev, rule, msg, _) in enumerate(issues, 1)]
        if rv.get("note") and reviewed:
            lines.append(f"- Review note: {rv['note']}")
        lines.append("")
    contact_sheet(out, slides, status)
    lines += ["", CHECKLIST]
    (out / "report.md").write_text("\n".join(lines), "utf-8")
    review_path.write_text(json.dumps(review, indent=1), "utf-8")
    print(f"{deck.name}: {len(slides)} slides, {fails} FAIL. Folder: {out}")
    for ln in lines[2:]:
        if ln.startswith(("## ", "1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.")) and "clean · pass" not in ln:
            print(ln)
    todo = ",".join(map(str, unreviewed))
    print(f"\nNeeds visual review on the current layout: {todo or 'none'}" + ("\n\n" + CHECKLIST if todo else ""))
    return 0 if not fails and not unreviewed else 1


def mark(deck: Path, spec: str, verdict: str, note: str = "") -> int:
    path = deck.parent / "_check" / deck.stem / "review.json"
    review = json.loads(path.read_text("utf-8"))
    for i in parse_range(spec):
        r = review[str(i)]
        r.update(hash=r["current"], verdict=verdict, note=note)
    path.write_text(json.dumps(review, indent=1), "utf-8")
    print(f"marked {spec} {verdict}")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")  # Windows console is cp1252
    a = sys.argv[1:]
    if len(a) >= 4 and a[1] == "--mark":
        sys.exit(mark(Path(a[0]).resolve(), a[2], a[3], a[4] if len(a) > 4 else ""))
    if len(a) == 1:
        sys.exit(run(Path(a[0]).resolve()))
    print(__doc__); sys.exit(2)
